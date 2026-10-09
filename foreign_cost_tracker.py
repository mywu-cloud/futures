# -*- coding: utf-8 -*-
"""
台指期 外資持倉成本 每日追蹤（期交所官方資料，兩種成本算法並列）
======================================================================
用法（repo 根目錄執行；GitHub Actions 每個交易日自動執行）：
  python foreign_cost_tracker.py                      # 每日例行
  python foreign_cost_tracker.py --backfill-days 730  # 首次回補（預設 365 天）

資料來源（臺灣期貨交易所）：
  三大法人-區分各期貨契約（TXF，外資及陸資）
  期貨每日交易行情（TX 一般盤：近月高低收、結算價、次月收盤）

輸出（data/ 目錄；CSV 為主資料，xlsx / json 由 CSV 重建；資料無變動時不改寫任何檔案）：
  data/taifex_raw.csv             期交所每日原始資料
  data/taifex_calc.csv            原始資料 + 兩種自算持倉成本與損益
  data/futures_foreign_cost.xlsx  對照總表 / 期交所原始 / 走勢圖 / 說明
  data/latest.json                最新一日摘要 + 近 120 日走勢（網站前端讀取）

算法一「結算重置法」（與玩股網相同算法，2026/10/08 驗證數值一致）：
  結算日：成本 =（近月最高 + 最低）/ 2 重置；未實現以次月收盤計
  非結算日：淨口數與前日同號且同向加碼 → 以近月收盤價加權平均；同號減碼 → 成本不變；
            翻空翻多 → 成本 =（最高 + 最低）/ 2；淨口數 0 → 成本 0
  已實現：減碼口數 ×（收盤 − 前日成本），每個合約月（結算日後）重新累計
算法二「連續移動平均法」：
  當日成交均價 p = 多空交易契約金額淨額(千元)*1000 / (多空交易口數淨額*200)
     （淨交易口數為 0 或偏離收盤 >15% 時改用結算價 / 收盤價）
  同向加碼加權平均；反向先沖銷計已實現，超出部分以 p 為新成本；不因結算重置

損益單位：萬元（每點 200 元）。本程式產出之數據僅供參考，不構成任何投資建議。
"""
import argparse
import csv
import io
import json
import os
import re
import sys
import time
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TZ = ZoneInfo("Asia/Taipei")  # Actions 執行環境為 UTC，一律以台北時間為準
TAIFEX_INST_URL = "https://www.taifex.com.tw/cht/3/futContractsDateDown"
TAIFEX_PRICE_URL = "https://www.taifex.com.tw/cht/3/futDataDown"
MULT = 200          # 大台每點 200 元
WAN = 1e4           # 損益單位：萬元
DATA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")
HISTORY_DAYS = 120  # latest.json 附帶的走勢天數

HEADERS = {"User-Agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                          "(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36"),
           "Accept-Language": "zh-TW,zh;q=0.9,en;q=0.8",
           "Referer": "https://www.taifex.com.tw/cht/3/futContractsDate"}

RAW_COLS = ["日期", "多空交易口數淨額", "多空交易契約金額淨額(千元)", "未平倉淨口數",
            "未平倉契約金額淨額(千元)", "近月契約", "近月最高", "近月最低", "近月收盤",
            "近月結算價", "次月收盤"]
CALC_COLS = ["結算日", "淨口數變動",
             "重置法成本", "重置法未實現(萬元)", "重置法已實現(萬元)", "重置法總損益(萬元)",
             "當日成交均價", "連續法成本", "連續法未實現(萬元)", "連續法已實現累計(萬元)", "連續法總損益(萬元)"]
TEXT_COLS = ("日期", "近月契約")


def now():
    return datetime.now(TZ)


def log(msg):
    print(f"[{now():%Y-%m-%d %H:%M:%S}] {msg}", flush=True)


def to_num(s):
    if s is None or isinstance(s, (int, float)):
        return s
    s = str(s).replace(",", "").replace("+", "").strip()
    if s in ("", "-", "--"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def norm_date(s):
    return datetime.strptime(str(s).strip().replace("-", "/"), "%Y/%m/%d").strftime("%Y/%m/%d")


def sgn(x):
    return (x > 0) - (x < 0)


# ---------------------------------------------------------------- HTTP
def http_post(url, data, retries=3):
    last = None
    for i in range(retries):
        try:
            r = requests.post(url, data=data, headers=HEADERS, timeout=30)
            if r.status_code == 200:
                return r.content
            last = f"HTTP {r.status_code}"
        except Exception as e:  # noqa
            last = repr(e)
        log(f"  重試 {i + 1}/{retries}：{url} → {last}")
        time.sleep(3 * (i + 1))
    raise RuntimeError(f"請求失敗：{url}（{last}）")


def decode(content):
    for enc in ("cp950", "big5", "utf-8-sig"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            pass
    return content.decode("cp950", errors="ignore")


def csv_rows(text):
    return [r for r in csv.reader(io.StringIO(text)) if any(c.strip() for c in r)]


def find_col(header, *keys):
    for i, h in enumerate(header):
        if all(k in h for k in keys):
            return i
    raise KeyError(f"找不到欄位 {keys}；表頭={header}")


def check_csv(text, first_col):
    """期交所查無資料時回傳 HTML 或空內容 → 視為無資料；其他非預期內容 → 報錯"""
    rows = csv_rows(text)
    if rows and first_col in rows[0][0]:
        return rows
    head = text[:300].lower()
    if not text.strip() or "<html" in head or "<!doctype" in head or "查無" in text:
        return []
    raise ValueError(f"期交所回傳格式非預期：{text[:200]!r}")


# ---------------------------------------------------------------- 期交所解析
def parse_taifex_inst(text):
    rows = check_csv(text, "日期")
    if not rows:
        return {}
    h = [c.strip() for c in rows[0]]
    i_d, i_id = find_col(h, "日期"), find_col(h, "身份別")
    i_tq, i_ta = find_col(h, "多空", "交易", "口數", "淨額"), find_col(h, "多空", "交易", "金額", "淨額")
    i_oq, i_oa = find_col(h, "多空", "未平倉", "口數", "淨額"), find_col(h, "多空", "未平倉", "金額", "淨額")
    out = {}
    for r in rows[1:]:
        if len(r) > i_oa and "外資" in r[i_id]:
            d = norm_date(r[i_d])
            out[d] = {"日期": d, "多空交易口數淨額": to_num(r[i_tq]), "多空交易契約金額淨額(千元)": to_num(r[i_ta]),
                      "未平倉淨口數": to_num(r[i_oq]), "未平倉契約金額淨額(千元)": to_num(r[i_oa])}
    return out


def parse_taifex_price(text):
    """TX 月契約、一般盤：近月（含結算日當天到期契約）高低收與結算價、次月收盤；排除週契約與價差"""
    rows = check_csv(text, "交易日期")
    if not rows:
        return {}
    h = [c.strip() for c in rows[0]]
    i_d, i_c, i_m = find_col(h, "交易日期"), find_col(h, "契約"), find_col(h, "到期月份")
    i_hi, i_lo = find_col(h, "最高價"), find_col(h, "最低價")
    i_close, i_settle = find_col(h, "收盤價"), find_col(h, "結算價")
    try:
        i_sess = find_col(h, "交易時段")
    except KeyError:
        i_sess = None
    by_day = {}
    for r in rows[1:]:
        if len(r) <= max(i_settle, i_close) or r[i_c].strip() != "TX":
            continue
        ym = r[i_m].strip()
        if not re.fullmatch(r"\d{6}", ym) or (i_sess is not None and "一般" not in r[i_sess]):
            continue
        if to_num(r[i_close]) is None:
            continue
        by_day.setdefault(norm_date(r[i_d]), []).append(
            (ym, to_num(r[i_hi]), to_num(r[i_lo]), to_num(r[i_close]), to_num(r[i_settle])))
    out = {}
    for d, lst in by_day.items():
        lst.sort()
        near = lst[0]
        out[d] = {"近月契約": near[0], "近月最高": near[1], "近月最低": near[2], "近月收盤": near[3],
                  "近月結算價": near[4], "次月收盤": lst[1][3] if len(lst) > 1 else None}
    return out


def fetch_taifex(start, end):
    inst, price, s = {}, {}, start
    while s <= end:
        e = min(s + timedelta(days=29), end)
        a, b = s.strftime("%Y/%m/%d"), e.strftime("%Y/%m/%d")
        log(f"  期交所 {a} ~ {b}")
        inst.update(parse_taifex_inst(decode(http_post(
            TAIFEX_INST_URL, {"queryStartDate": a, "queryEndDate": b, "commodityId": "TXF"}))))
        price.update(parse_taifex_price(decode(http_post(
            TAIFEX_PRICE_URL, {"down_type": "1", "commodity_id": "TX", "commodity_id2": "",
                               "queryStartDate": a, "queryEndDate": b}))))
        s = e + timedelta(days=1)
        time.sleep(1.5)
    empty = {k: None for k in RAW_COLS[5:]}
    return {d: {**row, **price.get(d, empty)} for d, row in inst.items()}


# ---------------------------------------------------------------- 自算持倉成本
def is_settlement(rows, i):
    """結算日：近月結算價為 0／空白（期交所到期日的標示），或隔一交易日近月契約已換月"""
    r = rows[i]
    if r.get("近月收盤") and not r.get("近月結算價"):
        return True
    if i + 1 < len(rows) and r.get("近月契約") and rows[i + 1].get("近月契約"):
        return rows[i + 1]["近月契約"] != r["近月契約"]
    return False


def compute_cost(tf):
    rows = sorted(tf.values(), key=lambda x: x["日期"])
    out = []
    prev = None                      # 前一筆有效的重置法紀錄
    P = C = None                     # 連續法：部位、成本
    realized_ma = 0.0
    for i, r in enumerate(rows):
        rec = {k: r.get(k) for k in RAW_COLS} | {k: None for k in CALC_COLS}
        oi, close, hi, lo = r.get("未平倉淨口數"), r.get("近月收盤"), r.get("近月最高"), r.get("近月最低")
        if oi is None or not close or not hi or not lo:
            out.append(rec)
            continue
        f = is_settlement(rows, i)
        settle = r.get("近月結算價") or close

        # ---- 算法一：結算重置法（與玩股網同）
        cost, unreal, realized, chg = (hi + lo) / 2, 0.0, 0.0, 0
        if prev is not None:
            chg = oi - prev["oi"]
            if not f:
                if oi == 0:
                    cost = 0.0
                elif oi > 0 and prev["oi"] > 0:
                    cost = (prev["cost"] * prev["oi"] + chg * close) / oi if chg > 0 else prev["cost"]
                elif oi < 0 and prev["oi"] < 0:
                    cost = (prev["cost"] * prev["oi"] + chg * close) / oi if chg < 0 else prev["cost"]
            mark = (r.get("次月收盤") or close) if f else close
            unreal = (mark - cost) * oi * MULT / WAN
            reducing = sgn(oi) * sgn(chg) < 0
            if f:
                realized = prev["realized"] + (cost - prev["cost"]) * oi * MULT / WAN
            elif prev["f"]:
                realized = (close - prev["cost"]) * -chg * MULT / WAN if reducing else 0.0
            else:
                realized = prev["realized"] + ((close - prev["cost"]) * -chg * MULT / WAN if reducing else 0.0)
        prev = {"oi": oi, "cost": cost, "realized": realized, "f": f}

        # ---- 算法二：連續移動平均法
        tq, ta = r.get("多空交易口數淨額"), r.get("多空交易契約金額淨額(千元)")
        p = settle
        if tq and ta is not None:
            p_try = ta * 1000 / (tq * MULT)
            if abs(p_try - close) / close < 0.15:
                p = p_try
        if P is None:
            P, C = oi, settle
        else:
            q = oi - P
            if q:
                if P == 0:
                    P, C = q, p
                elif (q > 0) == (P > 0):
                    C = (abs(P) * C + abs(q) * p) / (abs(P) + abs(q))
                    P += q
                else:
                    closed = min(abs(q), abs(P))
                    realized_ma += (p - C) * closed * sgn(P) * MULT
                    if abs(q) < abs(P):
                        P += q
                    elif abs(q) == abs(P):
                        P, C = 0, None
                    else:
                        P, C = P + q, p
        unreal_ma = (close - C) * P * MULT if C is not None else 0.0

        rec.update({"結算日": 1 if f else 0, "淨口數變動": chg,
                    "重置法成本": round(cost, 2), "重置法未實現(萬元)": round(unreal, 2),
                    "重置法已實現(萬元)": round(realized, 2), "重置法總損益(萬元)": round(unreal + realized, 2),
                    "當日成交均價": round(p, 2), "連續法成本": round(C, 2) if C is not None else None,
                    "連續法未實現(萬元)": round(unreal_ma / WAN, 2),
                    "連續法已實現累計(萬元)": round(realized_ma / WAN, 2),
                    "連續法總損益(萬元)": round((unreal_ma + realized_ma) / WAN, 2)})
        out.append(rec)
    return out


# ---------------------------------------------------------------- CSV（主資料，git diff 可讀）
def load_csv(path, cols):
    store = {}
    if os.path.exists(path):
        with open(path, encoding="utf-8-sig", newline="") as f:
            for rec in csv.DictReader(f):
                if rec.get("日期"):
                    store[rec["日期"]] = {c: (rec.get(c) or None) if c in TEXT_COLS else to_num(rec.get(c))
                                         for c in cols}
    return store


def clean(v):
    return int(v) if isinstance(v, float) and v.is_integer() else v


def csv_text(cols, rows):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(cols)
    for r in rows:
        w.writerow(["" if r.get(c) is None else clean(r.get(c)) for c in cols])
    return buf.getvalue()


def read_text(path):
    if not os.path.exists(path):
        return None
    with open(path, encoding="utf-8-sig", newline="") as f:
        return f.read()


def write_text(path, text):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        f.write(text)


# ---------------------------------------------------------------- Excel
FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F2937")
SETTLE_FILL = PatternFill("solid", fgColor="FEF3C7")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=10)
LINE = Border(bottom=Side(style="thin", color="D1D5DB"))
NUM_RG = '[Red]#,##0.00;[Color10]-#,##0.00;0'  # 台股慣例：正紅負綠
INT_RG = '[Red]#,##0;[Color10]-#,##0;0'
NUM2, INT0 = '#,##0.00', '#,##0'


def write_sheet(ws, cols, rows, fmts, widths=None):
    ws.append(cols)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 36
    for r in rows:
        ws.append([r.get(c) for c in cols])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.border = BODY_FONT, LINE
            if fmts.get(cols[c.column - 1]):
                c.number_format = fmts[cols[c.column - 1]]
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(c, max(11, min(len(c) * 1.6 + 2, 20)))
    ws.freeze_panes = "B2"


def build_workbook(path, calc, data_date):
    wb = Workbook()
    ws = wb.active
    ws.title = "對照總表"
    cols = ["日期", "結算日", "近月收盤", "未平倉淨口數", "淨口數變動",
            "重置法成本", "連續法成本", "成本差異(重置-連續)", "收盤距重置法成本",
            "重置法未實現(萬元)", "重置法已實現(萬元)", "重置法總損益(萬元)",
            "連續法未實現(萬元)", "連續法總損益(萬元)"]
    rows = [r | {"結算日": "結算" if r.get("結算日") else ""} for r in calc]
    write_sheet(ws, cols, rows, {
        "近月收盤": INT0, "未平倉淨口數": INT_RG, "淨口數變動": INT_RG,
        "重置法成本": NUM2, "連續法成本": NUM2, "成本差異(重置-連續)": NUM_RG, "收盤距重置法成本": NUM_RG,
        "重置法未實現(萬元)": NUM_RG, "重置法已實現(萬元)": NUM_RG, "重置法總損益(萬元)": NUM_RG,
        "連續法未實現(萬元)": NUM_RG, "連續法總損益(萬元)": NUM_RG}, {"日期": 12, "結算日": 8})
    for i, r in enumerate(rows, 2):  # 差異欄用公式，兩邊皆有值才計算
        ws[f"H{i}"] = f'=IF(AND(ISNUMBER(F{i}),ISNUMBER(G{i})),F{i}-G{i},"")'
        ws[f"I{i}"] = f'=IF(AND(ISNUMBER(C{i}),ISNUMBER(F{i})),C{i}-F{i},"")'
        if r["結算日"]:
            ws[f"B{i}"].fill = SETTLE_FILL

    write_sheet(wb.create_sheet("期交所原始"), RAW_COLS + CALC_COLS, calc, {
        "多空交易口數淨額": INT_RG, "多空交易契約金額淨額(千元)": INT_RG, "未平倉淨口數": INT_RG,
        "未平倉契約金額淨額(千元)": INT_RG, "近月最高": INT0, "近月最低": INT0, "近月收盤": INT0,
        "近月結算價": INT0, "次月收盤": INT0, "淨口數變動": INT_RG,
        "重置法成本": NUM2, "重置法未實現(萬元)": NUM_RG, "重置法已實現(萬元)": NUM_RG, "重置法總損益(萬元)": NUM_RG,
        "當日成交均價": NUM2, "連續法成本": NUM2, "連續法未實現(萬元)": NUM_RG,
        "連續法已實現累計(萬元)": NUM_RG, "連續法總損益(萬元)": NUM_RG})

    if len(rows) >= 2:
        n = len(rows) + 1
        ch = LineChart()
        ch.title, ch.height, ch.width, ch.y_axis.title = "外資持倉成本 vs 台指期近月收盤", 11, 28, "點"
        for col in (3, 6, 7):
            ch.add_data(Reference(ws, min_col=col, min_row=1, max_row=n), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
        vals = [v for r in rows for v in (r.get("近月收盤"), r.get("重置法成本"), r.get("連續法成本")) if v]
        if vals:
            ch.y_axis.scaling.min = int(min(vals) // 1000 * 1000)
            ch.y_axis.scaling.max = int(-(-max(vals) // 1000) * 1000)
        ch.x_axis.number_format, ch.x_axis.tickLblSkip = "@", max(1, len(rows) // 12)
        ch.x_axis.delete = ch.y_axis.delete = False  # 新版 Excel 預設會隱藏 openpyxl 圖表座標軸
        for s, color in zip(ch.series, ("6B7280", "DC2626", "2563EB")):
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width = 15000
            s.smooth = False
        wb.create_sheet("走勢圖").add_chart(ch, "B2")

    ws4 = wb.create_sheet("說明")
    for r in [("項目", "說明"),
              ("用途", "台指期外資持倉成本：以期交所官方資料自算兩種成本線並列對照，GitHub Actions 每交易日更新"),
              ("資料來源", "臺灣期貨交易所：三大法人-區分各期貨契約(TXF 外資及陸資)、期貨每日交易行情(TX 月契約一般盤)"),
              ("重置法", "與玩股網相同算法：結算日以近月(最高+最低)/2 重置成本、未實現以次月收盤計；"
                      "非結算日同向加碼以收盤價加權、減碼成本不變、翻向以(最高+最低)/2 重設；已實現每個合約月重新累計"),
              ("連續法", "移動平均成本法：成交均價=交易契約金額淨額/交易口數淨額/200（偏離收盤>15%改用結算價）；"
                      "反向沖銷計已實現、自回補起始日累計，不因結算重置"),
              ("結算日", "近月結算價為 0（到期日）或隔日近月換月者標為結算日，對照總表以黃底標示"),
              ("損益單位", "萬元（每點 200 元）"),
              ("色彩慣例", "台股慣例：正值/增加=紅，負值/減少=綠"),
              ("公式欄", "對照總表 H、I 欄為公式，兩邊皆有數值時才計算"),
              ("資料日期", data_date),
              ("免責聲明", "數據僅供參考，不構成任何投資建議；資料以期交所公告為準，使用者須遵守期交所資料使用規範")]:
        ws4.append(r)
    for c in ws4[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for row in ws4.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment = BODY_FONT, Alignment(wrap_text=True, vertical="top")
    ws4.column_dimensions["A"].width, ws4.column_dimensions["B"].width = 14, 100
    wb.save(path)
    log(f"已寫入 {path}")


def build_latest(calc, status):
    valid = [r for r in calc if r.get("重置法成本") is not None]
    last = valid[-1] if valid else None
    summary = None
    if last:
        summary = {
            "date": last["日期"], "close": clean(last["近月收盤"]), "contract": last["近月契約"],
            "net_oi": clean(last["未平倉淨口數"]), "net_oi_change": clean(last["淨口數變動"]),
            "settlement_day": bool(last["結算日"]),
            "reset": {"cost": last["重置法成本"], "unrealized": last["重置法未實現(萬元)"],
                      "realized": last["重置法已實現(萬元)"], "total": last["重置法總損益(萬元)"]},
            "continuous": {"cost": last["連續法成本"], "unrealized": last["連續法未實現(萬元)"],
                           "realized": last["連續法已實現累計(萬元)"], "total": last["連續法總損益(萬元)"]}}
    return {
        "updated_at": now().strftime("%Y-%m-%d %H:%M"), "status": status, "unit": "損益單位：萬元",
        "latest": summary,
        "history": {"fields": ["date", "close", "reset_cost", "continuous_cost", "net_oi"],
                    "rows": [[r["日期"], clean(r["近月收盤"]), r["重置法成本"], r["連續法成本"],
                              clean(r["未平倉淨口數"])] for r in valid[-HISTORY_DAYS:]]},
        "disclaimer": "數據僅供參考，不構成任何投資建議；資料來源：臺灣期貨交易所"}


# ---------------------------------------------------------------- 主流程
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data-dir", default=DATA_DIR)
    ap.add_argument("--backfill-days", type=int, default=365)
    ap.add_argument("--rebuild", action="store_true", help="CSV 未變動也強制重建 xlsx / json")
    a = ap.parse_args()

    os.makedirs(a.data_dir, exist_ok=True)
    P = lambda n: os.path.join(a.data_dir, n)  # noqa
    tf = load_csv(P("taifex_raw.csv"), RAW_COLS)
    log(f"既有資料：期交所 {len(tf)} 筆")
    status = "ok"

    try:
        today = now().date()
        start = (datetime.strptime(max(tf), "%Y/%m/%d").date() - timedelta(days=3)) if tf \
            else today - timedelta(days=a.backfill_days)
        new = fetch_taifex(start, today)
        if not new and (today - start).days >= 14:
            raise RuntimeError(f"{start} ~ {today} 期間查無任何資料，可能被期交所阻擋或網站改版")
        tf.update(new)
        log(f"期交所新增/更新 {len(new)} 筆")
    except Exception as e:
        status = f"error: {e}"
        log(f"期交所抓取失敗：{e}")

    if not tf:
        log("沒有任何資料，結束")
        sys.exit(1)

    calc = compute_cost(tf)
    raw_txt = csv_text(RAW_COLS, [tf[d] for d in sorted(tf)])
    calc_txt = csv_text(RAW_COLS + CALC_COLS, calc)
    unchanged = (raw_txt == read_text(P("taifex_raw.csv")) and calc_txt == read_text(P("taifex_calc.csv"))
                 and os.path.exists(P("futures_foreign_cost.xlsx")) and os.path.exists(P("latest.json")))

    if unchanged and status == "ok" and not a.rebuild:
        log("資料無變動（假日或尚未公布），不改寫檔案")
    else:
        write_text(P("taifex_raw.csv"), raw_txt)
        write_text(P("taifex_calc.csv"), calc_txt)
        build_workbook(P("futures_foreign_cost.xlsx"), calc, max(tf))
        with open(P("latest.json"), "w", encoding="utf-8") as f:
            json.dump(build_latest(calc, {"taifex": status}), f, ensure_ascii=False, indent=1)
        last = next((r for r in reversed(calc) if r.get("重置法成本") is not None), None)
        if last:
            log(f"{last['日期']} 收盤 {last['近月收盤']:,.0f}｜重置法成本 {last['重置法成本']:,.2f}｜"
                f"連續法成本 {last['連續法成本']}｜淨口數 {last['未平倉淨口數']:,.0f}")

    sys.exit(1 if status != "ok" else 0)  # 失敗 → Actions 標紅並寄信（已取得的資料仍會 commit）


if __name__ == "__main__":
    main()
